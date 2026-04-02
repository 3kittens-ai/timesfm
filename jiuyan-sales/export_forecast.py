import calendar
import argparse
import json
import math
import os
import sqlite3
import sys
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

BASE_DIR = Path(__file__).resolve().parent
REPO_ROOT = BASE_DIR.parent
JIUYAN_ROOT = REPO_ROOT.parent / "jiuyan"

# --- Configuration ---
DB_PATH = JIUYAN_ROOT / "sales_filtered_database" / "sales_filtered.sqlite"
OUTPUT_DIR = BASE_DIR / "outputs"
FORECAST_JSON_PATH = OUTPUT_DIR / "jiuyan_forecasts.json"
TARGET_CATEGORIES = ["线组", "鱼钩", "加长子线", "无结子线"]
HISTORY_START_YEAR = 2024
INVENTORY_PLAN_TURNOVER_DAYS = 45
PRODUCTION_LOT_SIZE = 50
DEFAULT_FORECAST_MONTHS = 12
DEMAND_SHEET_NAME = "AI 需求预测"
BOUNDS_SHEET_NAME = "AI 需求预测上下限参考"


def ensure_inventory_snapshot_column(conn):
    existing_columns = {
        row[1]
        for row in conn.execute("PRAGMA table_info(dim_sku)").fetchall()
    }
    if "inventory_snapshot_at" not in existing_columns:
        conn.execute("ALTER TABLE dim_sku ADD COLUMN inventory_snapshot_at TEXT")
        conn.commit()


def read_skus_from_file(filepath):
    file_path = Path(filepath)
    if not file_path.exists():
        raise FileNotFoundError(f"SKU 范围文件不存在: {filepath}")

    if file_path.suffix.lower() in {".xlsx", ".xls"}:
        df_scope = pd.read_excel(file_path)
    else:
        df_scope = pd.read_csv(file_path)

    normalized_map = {
        str(column).strip().lower().replace("_", "").replace(" ", ""): column
        for column in df_scope.columns
    }
    candidate_keys = ["商品编码", "sku_code", "sku", "barcode", "编码"]

    source_col = None
    for key in candidate_keys:
        normalized_key = key.strip().lower().replace("_", "").replace(" ", "")
        if normalized_key in normalized_map:
            source_col = normalized_map[normalized_key]
            break

    if source_col is None:
        raise ValueError("SKU 范围文件中未找到商品编码列")

    result = df_scope[[source_col]].copy()
    result.columns = ["sku_code"]
    result["sku_code"] = result["sku_code"].astype(str).str.strip()
    result = result[result["sku_code"] != ""]
    return result.drop_duplicates(subset=["sku_code"])


def load_forecast_json(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        raw_data = json.load(f)

    if not isinstance(raw_data, dict) or not raw_data:
        raise ValueError(f"预测文件为空或格式不正确: {filepath}")

    display_forecast_keys = None
    forecast_map = {}

    for sku_code, payload in raw_data.items():
        if not isinstance(payload, dict):
            continue

        months = payload.get("months") or []
        forecast_values = payload.get("forecast") or []
        lower_80_values = payload.get("lower_80") or []
        upper_80_values = payload.get("upper_80") or []

        if len(months) != len(forecast_values):
            raise ValueError(f"SKU {sku_code} 的 months 与 forecast 长度不一致")
        if lower_80_values and len(months) != len(lower_80_values):
            raise ValueError(f"SKU {sku_code} 的 months 与 lower_80 长度不一致")
        if upper_80_values and len(months) != len(upper_80_values):
            raise ValueError(f"SKU {sku_code} 的 months 与 upper_80 长度不一致")

        if display_forecast_keys is None:
            display_forecast_keys = list(months)

        forecast_map[str(sku_code)] = {
            str(month): {
                "forecast": float(forecast or 0),
                "lower_80": float(lower_80 or 0),
                "upper_80": float(upper_80 or 0),
            }
            for month, forecast, lower_80, upper_80 in zip(
                months,
                forecast_values,
                lower_80_values or [0] * len(months),
                upper_80_values or [0] * len(months),
            )
        }

    if not display_forecast_keys:
        raise ValueError(f"预测文件中没有可用 months: {filepath}")

    return display_forecast_keys, forecast_map

def inject_cached_values_for_numbers(filepath, sheet_formula_caches):
    """
    Inject cached values into the generated Excel file's XML.
    This fixes an issue where Apple Numbers evaluates openpyxl-generated formulas to 0.
    sheet_formula_caches: dict of {sheet_index: {cell_ref: cached_value}}
    """
    temp_path = filepath + ".tmp"
    with zipfile.ZipFile(filepath, 'r') as zin, zipfile.ZipFile(temp_path, 'w') as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            modified = False
            
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                try:
                    sheet_index = int(item.filename.replace("xl/worksheets/sheet", "").replace(".xml", ""))
                except ValueError:
                    sheet_index = -1
                
                if sheet_index in sheet_formula_caches:
                    caches = sheet_formula_caches[sheet_index]
                    parser = ET.XMLParser(target=ET.TreeBuilder(insert_comments=True))
                    root = ET.fromstring(content, parser=parser)
                    
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    ET.register_namespace('', ns['s'])
                    
                    for row in root.findall('.//s:row', ns):
                        for c in row.findall('.//s:c', ns):
                            ref = c.get('r')
                            if ref in caches:
                                v_val = caches[ref]
                                f_tag = c.find('s:f', ns)
                                if f_tag is not None:
                                    v_tag = c.find('s:v', ns)
                                    if v_tag is None:
                                        v_tag = ET.SubElement(c, 'v')
                                    v_tag.text = str(v_val)
                                    if isinstance(v_val, str):
                                        c.set('t', 'str')
                                    modified = True
                    if modified:
                        content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            zout.writestr(item, content)
    os.replace(temp_path, filepath)

def month_key(year, month):
    return f"{year}-{month:02d}"




def add_months(year, month, offset):
    total = (year * 12 + (month - 1)) + offset
    new_year = total // 12
    new_month = total % 12 + 1
    return new_year, new_month


def get_prev_ym(year, month):
    return add_months(year, month, -1)


def get_days_in_month(year, month):
    return calendar.monthrange(year, month)[1]


def safe_div_pct(current_value, previous_value):
    try:
        if previous_value in (None, 0) or pd.isna(previous_value):
            return None
        return (current_value / previous_value - 1) * 100
    except Exception:
        return None


def ceil_number(value):
    if value is None or value == "" or pd.isna(value):
        return 0
    return math.ceil(float(value))


def format_pct(value):
    if value is None or pd.isna(value):
        return ""
    return f"{math.ceil(float(value))}%"


def ceil_to_multiple(value, multiple):
    if value <= 0:
        return 0
    return int(math.ceil(float(value) / multiple) * multiple)


def calc_turnover_display(quantity, periods, fallback_month):
    if quantity is None or pd.isna(quantity) or quantity <= 0:
        return 0

    remaining_qty = float(quantity)
    consumed_days = 0.0
    total_period_days = sum(period["days"] for period in periods)

    for period in periods:
        period_days = period["days"]
        daily_rate = period["daily_rate"]

        if period_days <= 0:
            continue

        if daily_rate <= 0:
            consumed_days += period_days
            continue

        period_demand = daily_rate * period_days
        if remaining_qty > period_demand:
            remaining_qty -= period_demand
            consumed_days += period_days
            continue

        consumed_days += remaining_qty / daily_rate
        return math.ceil(consumed_days)

    if consumed_days > total_period_days:
        consumed_days = total_period_days

    if remaining_qty > 0:
        return f"{fallback_month}月底以后"

    return math.ceil(consumed_days)


def quantity_for_target_days(target_days, periods):
    remaining_days = float(target_days)
    required_qty = 0.0

    for period in periods:
        if remaining_days <= 0:
            break
        days_to_take = min(remaining_days, period["days"])
        required_qty += period["daily_rate"] * days_to_take
        remaining_days -= days_to_take

    if remaining_days > 0 and periods:
        required_qty += periods[-1]["daily_rate"] * remaining_days

    return required_qty


def build_turnover_formula(quantity_ref, current_demand_ref, current_daily_expr, future_demand_refs, future_daily_refs, period_days, last_labels):
    demand_terms = [current_demand_ref] + future_demand_refs
    rate_terms = [current_daily_expr] + future_daily_refs
    cumulative_demands = []
    running_demand = ""
    for idx, demand_ref in enumerate(demand_terms):
        running_demand = demand_ref if idx == 0 else f"({running_demand}+{demand_ref})"
        cumulative_demands.append(running_demand)

    elapsed_days = []
    running_days = 0
    for idx, days in enumerate(period_days):
        elapsed_days.append(str(running_days))
        running_days += days

    nested_expr = f'"{last_labels[-1]}"'
    for idx in range(len(demand_terms) - 1, -1, -1):
        rate_ref = rate_terms[idx]
        demand_cum_ref = cumulative_demands[idx]
        elapsed_ref = elapsed_days[idx]
        if idx == 0:
            nested_expr = (
                f'IF({quantity_ref}<={demand_cum_ref},'
                f'IF({rate_ref}=0,0,ROUNDUP({quantity_ref}/{rate_ref},0)),'
                f'{nested_expr})'
            )
        else:
            prev_cum_ref = cumulative_demands[idx - 1]
            nested_expr = (
                f'IF({quantity_ref}<={demand_cum_ref},'
                f'IF({rate_ref}=0,{elapsed_ref},ROUNDUP({elapsed_ref}+({quantity_ref}-{prev_cum_ref})/{rate_ref},0)),'
                f'{nested_expr})'
            )

    return f'=IF({quantity_ref}<=0,0,{nested_expr})'


def build_inventory_plan_formula(total_quantity_ref, current_daily_expr, future_daily_refs, period_days):
    target_days_ref = "'常量'!$B$4"
    rate_terms = [current_daily_expr] + future_daily_refs
    take_days_terms = []
    elapsed_days = 0
    for days in period_days:
        take_days_terms.append(f"MAX(MIN({target_days_ref}-{elapsed_days},{days}),0)")
        elapsed_days += days

    need_qty_expr = "+".join(
        f"({take_days_terms[idx]})*({rate_terms[idx]})" for idx in range(len(rate_terms))
    )
    extra_days_expr = f"MAX({target_days_ref}-{sum(period_days)},0)"
    raw_gap_expr = f"MAX(({need_qty_expr})+({extra_days_expr})*({rate_terms[-1]})-({total_quantity_ref}),0)"
    return f'=IF(({raw_gap_expr})<=0,0,ROUNDUP(({raw_gap_expr})/{PRODUCTION_LOT_SIZE},0)*{PRODUCTION_LOT_SIZE})'


def export_forecast(
    sku_scope_file=None,
    cutoff_date_str=None,
    output_prefix="ai-forecast",
    forecast_json_path=FORECAST_JSON_PATH,
    forecast_months=DEFAULT_FORECAST_MONTHS,
):
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    display_forecast_keys, forecast_json_map = load_forecast_json(forecast_json_path)
    max_forecast_months = len(display_forecast_keys)
    effective_forecast_months = max(1, min(int(forecast_months), max_forecast_months))
    display_forecast_keys = display_forecast_keys[:effective_forecast_months]

    conn = sqlite3.connect(DB_PATH)
    ensure_inventory_snapshot_column(conn)

    if cutoff_date_str:
        latest_date_str = pd.read_sql(
            "SELECT MAX(sale_date) AS latest_date FROM sales WHERE sale_date <= ?",
            conn,
            params=(cutoff_date_str,),
        ).iloc[0, 0]
    else:
        latest_date_str = pd.read_sql("SELECT MAX(sale_date) AS latest_date FROM sales", conn).iloc[0, 0]
    if not latest_date_str:
        conn.close()
        print("❌ 数据库中没有销售数据项。")
        return None

    latest_date = datetime.strptime(latest_date_str, "%Y-%m-%d")
    base_year = latest_date.year
    base_month = latest_date.month
    base_month_key = month_key(base_year, base_month)
    inventory_latest_update = pd.read_sql(
        "SELECT MAX(inventory_snapshot_at) AS updated_at FROM dim_sku",
        conn,
    ).iloc[0, 0]
    inventory_latest_update_display = str(inventory_latest_update).split(" ")[0] if inventory_latest_update else ""
    inventory_snapshot_note = ""
    if inventory_latest_update_display and latest_date_str:
        inventory_snapshot_note = (
            f"生产计划中的当前库存数量已按照 {inventory_latest_update_display} 更新的库存数量，"
            f"扣减了 {inventory_latest_update_display} 之后至 {latest_date_str} 的实际销量。"
        )

    history_keys = []
    year_cursor, month_cursor = HISTORY_START_YEAR, 1
    while (year_cursor < base_year) or (year_cursor == base_year and month_cursor <= base_month):
        history_keys.append(month_key(year_cursor, month_cursor))
        year_cursor, month_cursor = add_months(year_cursor, month_cursor, 1)

    forecast_mom_keys = display_forecast_keys[1:]

    current_year_actual_month_cols = [month_key(base_year, m) for m in range(1, base_month)]
    prev_year = base_year - 1
    prev2_year = base_year - 2
    prev_y1_cols = [month_key(prev_year, m) for m in range(1, 13)]
    prev_y2_cols = [month_key(prev2_year, m) for m in range(1, 13)]

    current_month_mean_col = f"{base_month} 月预测（本月均值参考）"
    current_month_window_col = f"{base_month} 月预测（7-15日参考）"
    current_month_ly_col = f"{base_month} 月预测（历史同期环比参考）"

    df_skus = pd.read_sql(
        f"""
        SELECT sku_code, sku_name, product_category, family_tags, supplier, creation_date,
               finished_inventory, purchase_in_transit, inventory_snapshot_at
        FROM dim_sku
        WHERE product_category IN ({",".join(["?"] * len(TARGET_CATEGORIES))})
        """,
        conn,
        params=TARGET_CATEGORIES,
    )
    df_skus["family_tags"] = df_skus["family_tags"].fillna("")
    df_skus["supplier"] = df_skus["supplier"].fillna("")
    df_skus["creation_date"] = df_skus["creation_date"].fillna("")
    df_skus["finished_inventory"] = df_skus["finished_inventory"].fillna(0)
    df_skus["purchase_in_transit"] = df_skus["purchase_in_transit"].fillna(0)
    df_skus["inventory_snapshot_at"] = df_skus["inventory_snapshot_at"].fillna("")

    if sku_scope_file:
        df_scope = read_skus_from_file(sku_scope_file)
        scope_codes = set(df_scope["sku_code"].astype(str))
        df_skus = df_skus[df_skus["sku_code"].astype(str).isin(scope_codes)].copy()

    if not df_skus.empty:
        sku_codes = df_skus["sku_code"].astype(str).tolist()
        placeholders = ",".join(["?"] * len(sku_codes))
        df_post_snapshot_sales = pd.read_sql(
            f"""
            SELECT d.sku_code,
                   COALESCE(SUM(s.sales_volume), 0) AS post_snapshot_sales_volume
            FROM dim_sku d
            LEFT JOIN sales s
             ON s.barcode = d.sku_code
             AND d.inventory_snapshot_at IS NOT NULL
             AND TRIM(d.inventory_snapshot_at) != ''
             AND s.sale_date > substr(d.inventory_snapshot_at, 1, 10)
             AND s.sale_date <= ?
            WHERE d.sku_code IN ({placeholders})
            GROUP BY d.sku_code
            """,
            conn,
            params=[latest_date_str, *sku_codes],
        )
        post_snapshot_sales_map = {
            str(row["sku_code"]): float(row["post_snapshot_sales_volume"] or 0)
            for _, row in df_post_snapshot_sales.iterrows()
        }
    else:
        post_snapshot_sales_map = {}

    df_skus["post_snapshot_sales_volume"] = df_skus["sku_code"].astype(str).map(post_snapshot_sales_map).fillna(0)
    df_skus["effective_finished_inventory"] = (
        df_skus["finished_inventory"].astype(float) - df_skus["post_snapshot_sales_volume"].astype(float)
    )

    included_categories = "、".join(sorted(df_skus["product_category"].dropna().astype(str).unique().tolist()))
    included_sku_count = int(df_skus["sku_code"].nunique())

    df_monthly = pd.read_sql(
        """
        SELECT substr(sale_date, 1, 7) AS ym, barcode, SUM(sales_volume) AS qty
        FROM sales
        WHERE sale_date >= ?
          AND sale_date <= ?
        GROUP BY ym, barcode
        """,
        conn,
        params=(f"{HISTORY_START_YEAR}-01-01", latest_date_str),
    )
    pivot_monthly = df_monthly.pivot(index="barcode", columns="ym", values="qty").fillna(0)
    pivot_dict = pivot_monthly.to_dict(orient="index")

    def get_actual_qty(code, ym_key):
        return float(pivot_dict.get(code, {}).get(ym_key, 0.0))

    def get_group_actual_qty(group_rows, ym_key):
        return sum(get_actual_qty(r["商品编码"], ym_key) for r in group_rows if str(r["商品编码"]) not in {"合计", "环比", "同比"})

    lookback_months = [month_key(*add_months(base_year, base_month, -i)) for i in range(1, 13)]

    start_window = (latest_date - timedelta(days=14)).strftime("%Y-%m-%d")
    df_daily = pd.read_sql(
        """
        SELECT sale_date, barcode, SUM(sales_volume) AS qty
        FROM sales
        WHERE sale_date >= ? AND sale_date <= ?
        GROUP BY sale_date, barcode
        """,
        conn,
        params=(start_window, latest_date_str),
    )
    if not df_daily.empty:
        df_daily["sale_date_dt"] = pd.to_datetime(df_daily["sale_date"])
        daily_7 = df_daily[df_daily["sale_date_dt"] >= (latest_date - timedelta(days=6))].groupby("barcode")["qty"].sum()
        daily_15 = df_daily.groupby("barcode")["qty"].sum()
    else:
        daily_7 = pd.Series(dtype="float64")
        daily_15 = pd.Series(dtype="float64")

    conn.close()

    days_in_base_month = get_days_in_month(base_year, base_month)
    days_elapsed = latest_date.day
    remaining_days_in_base_month = max(days_in_base_month - days_elapsed, 0)
    fallback_month = int(display_forecast_keys[-1].split("-")[1])

    current_remaining_demand_col = f"{base_month}月剩余需求"
    future_demand_cols = [f"{int(ym_key.split('-')[1])}月需求" for ym_key in display_forecast_keys[1:]]
    future_daily_cols = [f"{int(ym_key.split('-')[1])}月每日用量" for ym_key in display_forecast_keys[1:]]
    current_gap_col = f"{base_month}月销售差额"
    next_month_num = int(display_forecast_keys[1].split("-")[1]) if len(display_forecast_keys) > 1 else None
    cumulative_gap_col = f"{base_month}月+{next_month_num}月销售差额" if next_month_num else None

    all_results = []
    plan_results = []
    bounds_results = []
    grouped = df_skus.groupby(["product_category", "family_tags"], dropna=False)

    for (category, tags), group in grouped:
        family_name = f"[{category}] {tags if tags else ''}".strip()
        group_rows = []
        group_bounds_rows = []

        for _, sku in group.iterrows():
            code = sku["sku_code"]
            row = {
                "商品编码": code,
                "商品名称": sku["sku_name"],
                "商品分类和标签": family_name,
            }

            for ym_key in history_keys:
                row[ym_key] = get_actual_qty(code, ym_key)

            row[f"{base_year}年实销合计"] = sum(row.get(month_key(base_year, m), 0) for m in range(1, base_month))
            row[f"{prev_year}年实销合计"] = sum(row.get(month_key(prev_year, m), 0) for m in range(1, 13))
            row[f"{prev2_year}年实销合计"] = sum(row.get(month_key(prev2_year, m), 0) for m in range(1, 13))

            lookback_values = [get_actual_qty(code, ym_key) for ym_key in lookback_months]
            row["最近12个月最高(不含本月)"] = max(lookback_values) if lookback_values else 0
            row["最近12个月平均(不含本月)"] = (sum(lookback_values) / len(lookback_values)) if lookback_values else 0

            base_actual = row.get(base_month_key, 0)
            mean_forecast = (base_actual / days_elapsed) * days_in_base_month if days_elapsed else 0
            window_forecast = max(
                (daily_7.get(code, 0) / 7) * days_in_base_month,
                (daily_15.get(code, 0) / 15) * days_in_base_month,
            )

            prev_year_for_base, prev_month_for_base = get_prev_ym(base_year, base_month)
            prev_month_key = month_key(prev_year_for_base, prev_month_for_base)
            prev_actual = get_actual_qty(code, prev_month_key)
            ly_prev_same_month = get_actual_qty(code, month_key(base_year - 1, prev_month_for_base))
            ly_prev_prev_year, ly_prev_prev_month = get_prev_ym(base_year - 1, prev_month_for_base)
            ly_prev_prev_month_qty = get_actual_qty(code, month_key(ly_prev_prev_year, ly_prev_prev_month))
            ly_forecast = prev_actual * (
                ly_prev_prev_month_qty / ly_prev_same_month if ly_prev_same_month > 0 else 1.0
            )

            row[current_month_mean_col] = mean_forecast
            row[current_month_window_col] = window_forecast
            row[current_month_ly_col] = ly_forecast

            sku_forecasts = forecast_json_map.get(str(code), {})
            for forecast_key in display_forecast_keys:
                row[forecast_key] = sku_forecasts.get(forecast_key, {}).get("forecast", 0.0)

            row["预测合计"] = sum(row.get(ym_key, 0) for ym_key in display_forecast_keys)

            for mom_key in forecast_mom_keys:
                mom_year, mom_month = map(int, mom_key.split("-"))
                prev_mom_year, prev_mom_month = get_prev_ym(mom_year, mom_month)
                prev_mom_key = month_key(prev_mom_year, prev_mom_month)
                row[f"{mom_key}环比(%)"] = safe_div_pct(row.get(mom_key, 0), row.get(prev_mom_key, 0))

            on_hand_inventory = float(sku["effective_finished_inventory"] or 0)
            in_transit_inventory = float(sku["purchase_in_transit"] or 0)
            inventory_plus_transit = on_hand_inventory + in_transit_inventory
            current_remaining_demand = math.ceil(
                row[base_month_key] / days_in_base_month * remaining_days_in_base_month if days_in_base_month else 0
            )
            future_demands = [math.ceil(row.get(ym_key, 0)) for ym_key in display_forecast_keys[1:]]
            future_daily_usage = [
                row.get(ym_key, 0) / get_days_in_month(*map(int, ym_key.split("-")))
                for ym_key in display_forecast_keys[1:]
            ]

            periods = [
                {
                    "label": current_remaining_demand_col,
                    "days": remaining_days_in_base_month,
                    "daily_rate": row[base_month_key] / days_in_base_month if days_in_base_month else 0,
                }
            ]
            for ym_key in display_forecast_keys[1:]:
                forecast_year, forecast_month = map(int, ym_key.split("-"))
                periods.append(
                    {
                        "label": ym_key,
                        "days": get_days_in_month(forecast_year, forecast_month),
                        "daily_rate": row.get(ym_key, 0) / get_days_in_month(forecast_year, forecast_month),
                    }
                )

            inventory_turnover = calc_turnover_display(on_hand_inventory, periods, fallback_month)
            inventory_plus_transit_turnover = calc_turnover_display(inventory_plus_transit, periods, fallback_month)
            inventory_plus_transit_meets = (
                "是"
                if isinstance(inventory_plus_transit_turnover, str)
                or inventory_plus_transit_turnover >= INVENTORY_PLAN_TURNOVER_DAYS
                else "否"
            )
            qty_for_target_days = quantity_for_target_days(INVENTORY_PLAN_TURNOVER_DAYS, periods)
            inventory_plan_qty = ceil_to_multiple(max(qty_for_target_days - inventory_plus_transit, 0), PRODUCTION_LOT_SIZE)

            plan_row = {
                "商品编码": code,
                "商品名称": sku["sku_name"],
                "供应商": sku["supplier"],
                current_remaining_demand_col: current_remaining_demand,
                current_gap_col: max(math.ceil(current_remaining_demand - inventory_plus_transit), 0),
                "当前在途数量": in_transit_inventory,
                "当前库存数量": on_hand_inventory,
                "库存可周转天数": inventory_turnover,
                "库存+在途可周转天数": inventory_plus_transit_turnover,
                "库存+在途是否满足需求": inventory_plus_transit_meets,
                "库存计划量": inventory_plan_qty,
                "商品创建日期": sku["creation_date"],
            }
            for col_name, value in zip(future_demand_cols, future_demands):
                plan_row[col_name] = value
            for col_name, value in zip(future_daily_cols, future_daily_usage):
                plan_row[col_name] = value
            if cumulative_gap_col:
                plan_row[cumulative_gap_col] = max(
                    math.ceil(current_remaining_demand + (future_demands[0] if future_demands else 0) - inventory_plus_transit),
                    0,
                )

            bounds_row = {
                "商品编码": code,
                "商品名称": sku["sku_name"],
            }
            for forecast_key in display_forecast_keys:
                month_values = sku_forecasts.get(forecast_key, {})
                bounds_row[f"{forecast_key}__forecast"] = float(month_values.get("forecast", 0.0))
                bounds_row[f"{forecast_key}__lower_80"] = float(month_values.get("lower_80", 0.0))
                bounds_row[f"{forecast_key}__upper_80"] = float(month_values.get("upper_80", 0.0))

            group_rows.append(row)
            plan_results.append(plan_row)
            group_bounds_rows.append(bounds_row)

        subtotal = {
            "商品编码": "合计",
            "商品名称": family_name,
            "商品分类和标签": family_name,
            "is_total_row": True,
        }
        sum_targets = list(
            dict.fromkeys(
                history_keys
                + display_forecast_keys
                + [
                    "预测合计",
                    current_month_mean_col,
                    current_month_window_col,
                    current_month_ly_col,
                    "最近12个月最高(不含本月)",
                    "最近12个月平均(不含本月)",
                    f"{base_year}年实销合计",
                    f"{prev_year}年实销合计",
                    f"{prev2_year}年实销合计",
                ]
            )
        )
        for column in sum_targets:
            subtotal[column] = sum(r.get(column, 0) for r in group_rows)

        month_columns_for_pct_rows = list(
            dict.fromkeys(history_keys + display_forecast_keys)
        )
        mom_row = {
            "商品编码": "环比",
            "商品名称": family_name,
            "商品分类和标签": family_name,
            "is_mom_row": True,
        }
        yoy_row = {
            "商品编码": "同比",
            "商品名称": family_name,
            "商品分类和标签": family_name,
            "is_yoy_row": True,
        }

        for ym_key in month_columns_for_pct_rows:
            year, month = map(int, ym_key.split("-"))
            prev_year_value, prev_month_value = get_prev_ym(year, month)
            prev_key = month_key(prev_year_value, prev_month_value)
            current_value = subtotal.get(ym_key, 0)
            previous_value = subtotal.get(prev_key, get_group_actual_qty(group_rows, prev_key))
            mom_row[ym_key] = safe_div_pct(current_value, previous_value)

            last_year_key = month_key(year - 1, month)
            yoy_base_value = subtotal.get(last_year_key, get_group_actual_qty(group_rows, last_year_key))
            yoy_row[ym_key] = safe_div_pct(current_value, yoy_base_value)

        all_results.extend(group_rows)
        all_results.append(subtotal)
        all_results.append(mom_row)
        all_results.append(yoy_row)

        bounds_subtotal = {
            "商品编码": "合计",
            "商品名称": family_name,
            "is_total_row": True,
        }
        bounds_mom_row = {
            "商品编码": "环比",
            "商品名称": family_name,
            "is_mom_row": True,
        }
        bounds_yoy_row = {
            "商品编码": "同比",
            "商品名称": family_name,
            "is_yoy_row": True,
        }

        for forecast_key in display_forecast_keys:
            forecast_year, forecast_month = map(int, forecast_key.split("-"))
            prev_forecast_year, prev_forecast_month = get_prev_ym(forecast_year, forecast_month)
            prev_forecast_key = month_key(prev_forecast_year, prev_forecast_month)
            prev_actual_total = get_group_actual_qty(group_rows, prev_forecast_key)
            yoy_actual_total = get_group_actual_qty(group_rows, month_key(forecast_year - 1, forecast_month))

            for metric_key in ("forecast", "lower_80", "upper_80"):
                column_key = f"{forecast_key}__{metric_key}"
                bounds_subtotal[column_key] = sum(r.get(column_key, 0) for r in group_bounds_rows)

                if forecast_key == display_forecast_keys[0]:
                    previous_value = prev_actual_total
                else:
                    previous_value = bounds_subtotal.get(f"{prev_forecast_key}__{metric_key}", 0)

                bounds_mom_row[column_key] = safe_div_pct(bounds_subtotal[column_key], previous_value)
                bounds_yoy_row[column_key] = safe_div_pct(bounds_subtotal[column_key], yoy_actual_total)

        bounds_results.extend(group_bounds_rows)
        bounds_results.append(bounds_subtotal)
        bounds_results.append(bounds_mom_row)
        bounds_results.append(bounds_yoy_row)

    df_raw = pd.DataFrame(all_results)

    forecast_mom_display_cols = [f"{ym_key}环比(%)" for ym_key in forecast_mom_keys]
    forecast_header_labels = {ym_key: f"{ym_key}\n（手动调整后会自动更新关联值）" for ym_key in display_forecast_keys}
    final_cols = [
        "商品编码",
        "商品名称",
        *display_forecast_keys,
        "预测合计",
        *forecast_mom_display_cols,
        "最近12个月最高(不含本月)",
        "最近12个月平均(不含本月)",
        current_month_mean_col,
        current_month_window_col,
        current_month_ly_col,
        f"{base_year}年实销合计",
        *current_year_actual_month_cols,
        f"{prev_year}年实销合计",
        *prev_y1_cols,
        f"{prev2_year}年实销合计",
        *prev_y2_cols,
        "商品分类和标签",
    ]
    final_cols = [column for column in final_cols if column in df_raw.columns]

    month_cols_in_output = set(display_forecast_keys + current_year_actual_month_cols + prev_y1_cols + prev_y2_cols)
    numeric_cols_in_output = set(
        month_cols_in_output
        | {
            "预测合计",
            "最近12个月最高(不含本月)",
            "最近12个月平均(不含本月)",
            current_month_mean_col,
            current_month_window_col,
            current_month_ly_col,
            f"{base_year}年实销合计",
            f"{prev_year}年实销合计",
            f"{prev2_year}年实销合计",
        }
    )
    pct_cols_in_output = set(forecast_mom_display_cols)

    def process_row_final(row):
        is_pct_row = row.get("is_mom_row", False) is True or row.get("is_yoy_row", False) is True
        is_total_row = row.get("is_total_row", False) is True

        for column in final_cols:
            raw_value = row.get(column)
            if column in ("商品编码", "商品名称", "商品分类和标签"):
                continue

            if is_pct_row:
                row[column] = format_pct(raw_value) if column in month_cols_in_output else ""
                continue

            if column in pct_cols_in_output:
                row[column] = "" if is_total_row else format_pct(raw_value)
            elif column in numeric_cols_in_output:
                row[column] = ceil_number(raw_value)
            else:
                row[column] = raw_value

        return row

    df_result = df_raw.copy().apply(process_row_final, axis=1)

    plan_cols = [
        "商品编码",
        "商品名称",
        "供应商",
        current_remaining_demand_col,
        *future_demand_cols,
        current_gap_col,
        *([cumulative_gap_col] if cumulative_gap_col else []),
        *future_daily_cols,
        "当前在途数量",
        "当前库存数量",
        "库存可周转天数",
        "库存+在途可周转天数",
        "库存+在途是否满足需求",
        "库存计划量",
        "商品创建日期",
    ]
    df_plan_raw = pd.DataFrame(plan_results)
    df_plan_raw = df_plan_raw[plan_cols]

    plan_numeric_cols = {
        current_remaining_demand_col,
        *future_demand_cols,
        current_gap_col,
        *([cumulative_gap_col] if cumulative_gap_col else []),
        *future_daily_cols,
        "当前在途数量",
        "当前库存数量",
        "库存计划量",
    }
    plan_turnover_cols = {"库存可周转天数", "库存+在途可周转天数"}

    def process_plan_row(row):
        for column in plan_cols:
            value = row.get(column)
            if column in plan_numeric_cols:
                row[column] = ceil_number(value)
            elif column in plan_turnover_cols:
                row[column] = value if isinstance(value, str) else ceil_number(value)
        return row

    df_plan = df_plan_raw.copy().apply(process_plan_row, axis=1)
    df_constants = pd.DataFrame(
        [
            ("销售数据截止日期：", latest_date_str, ""),
            ("库存快照时间：", inventory_latest_update_display, inventory_snapshot_note),
            ("库存计划量可周转（天）：", INVENTORY_PLAN_TURNOVER_DAYS, "手动调整后会自动更新关联值"),
            ("AI 预测周期（月）：", len(display_forecast_keys), f"默认 {DEFAULT_FORECAST_MONTHS} 个月，可在导出前配置，最多不超过 jiuyan_forecasts.json"),
            ("产品分类：", included_categories, ""),
            ("SKU 数量：", included_sku_count, ""),
        ],
        columns=["参数", "值", "说明"],
    )
    formula_rows = [
        (
            "预测月销量",
            "TimesFM 是 Google 做的一个“通用时间序列基础模型”，可以把它理解成时间序列领域里类似大语言模型的东西：先在大量不同类型的时间序列上预训练，再拿来直接做零样本或少样本预测。最直接影响预测的因素是：历史销量轨迹本身：趋势、季节性、波动、断崖、长尾销量。",
        ),
        (
            "当前月预测（本月均值参考）",
            "当前月最新实际总销量 / 当前月已销售天数 x 当前月总天数。",
        ),
        (
            "当前月预测（7-15日参考）",
            "取 max(最近 7 日销量 / 7, 最近 15 日销量 / 15) x 当前月总天数。",
        ),
        (
            "当前月预测（历史同期环比参考）",
            "当前月的上个月实际销量 x 去年该上个月的上个月实际销量 / 去年该上个月实际销量；若分母为 0，则按 1.0 处理。",
        ),
        (
            "采购在途是否满足需求",
            "生产计划 sheet 中“库存+在途是否满足需求”列：若“库存+在途可周转天数”大于等于常量 sheet 中“库存计划量可周转（天）”，显示“是”；否则显示“否”。若可周转到预测周期最后一个月月底以后，也显示“是”。",
        ),
        (
            "成品周转天数",
            "按生产计划 sheet 的需求周期逐日扣减当前库存数量：当前库存数量会先基于“最近一次库存快照数量 - 快照后累计销量”得到有效库存，再先扣当前月剩余需求对应日耗，随后按后续各月每日用量顺序扣减；若可支撑到预测周期最后一个月月底以后，则显示“X月底以后”。",
        ),
        (
            "在途+库存合计可周转天数",
            "按与“成品周转天数”相同的日耗逻辑，用“有效当前库存数量 + 当前在途数量”计算还能支撑的天数；若可支撑到预测周期最后一个月月底以后，则显示“X月底以后”。",
        ),
        (
            "库存计划量",
            "先按常量 sheet 中“库存计划量可周转（天）”计算目标覆盖天数对应的需求量，再减去“当前库存数量 + 当前在途数量”；若结果大于 0，则向上按每 50 取整，否则为 0。",
        ),
    ]
    df_formula = pd.DataFrame(formula_rows, columns=["字段", "说明"])

    file_path = os.path.join(OUTPUT_DIR, f"{output_prefix}-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df_result[final_cols].to_excel(writer, index=False, sheet_name=DEMAND_SHEET_NAME)
        df_plan.to_excel(writer, index=False, sheet_name="生产计划")
        df_constants.to_excel(writer, index=False, sheet_name="常量")
        df_formula.to_excel(writer, index=False, sheet_name="公式")

    wb = load_workbook(file_path)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    ws_bounds = wb.create_sheet(BOUNDS_SHEET_NAME, 1)

    def raw_demand_row(excel_row):
        return all_results[excel_row - 2]

    demand_caches = {}
    plan_caches = {}

    class OpenpyxlFormulaWriter:
        def __init__(self, worksheet, cache_dict):
            self.worksheet = worksheet
            self.cache_dict = cache_dict

        def write_formula(self, cell_ref, formula, _cell_format=None, cached_value=None):
            self.worksheet[cell_ref] = formula
            if cached_value is not None and cached_value != "":
                self.cache_dict[cell_ref] = cached_value

    ws_x_demand = OpenpyxlFormulaWriter(wb[DEMAND_SHEET_NAME], demand_caches)
    ws_x_plan = OpenpyxlFormulaWriter(wb["生产计划"], plan_caches)

    total_row_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    grey_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    blue_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    red_fill = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    bold_font = Font(bold=True)
    header_font = Font(bold=True, size=11)
    note_font = Font(size=8)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    bounds_subheaders = [
        ("forecast", "预测值"),
        ("lower_80", "预测下限"),
        ("upper_80", "预测上限"),
    ]
    ws_bounds.freeze_panes = ws_bounds["A3"]
    ws_bounds.row_dimensions[1].height = 30
    ws_bounds.row_dimensions[2].height = 26
    ws_bounds.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws_bounds.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    ws_bounds.cell(row=1, column=1, value="商品编码")
    ws_bounds.cell(row=1, column=2, value="商品名称")

    bounds_col_idx = 3
    for ym_key in display_forecast_keys:
        ws_bounds.merge_cells(start_row=1, start_column=bounds_col_idx, end_row=1, end_column=bounds_col_idx + 2)
        ws_bounds.cell(row=1, column=bounds_col_idx, value=ym_key)
        for offset, (_, subheader_label) in enumerate(bounds_subheaders):
            ws_bounds.cell(row=2, column=bounds_col_idx + offset, value=subheader_label)
        bounds_col_idx += len(bounds_subheaders)

    for row_idx in (1, 2):
        for col_idx in range(1, ws_bounds.max_column + 1):
            cell = ws_bounds.cell(row=row_idx, column=col_idx)
            cell.fill = blue_fill if row_idx == 1 else green_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = border

    for row_idx, bounds_row in enumerate(bounds_results, start=3):
        is_total_row = bounds_row.get("is_total_row", False) is True
        is_mom_row = bounds_row.get("is_mom_row", False) is True
        is_yoy_row = bounds_row.get("is_yoy_row", False) is True

        ws_bounds.row_dimensions[row_idx].height = 22
        ws_bounds.cell(row=row_idx, column=1, value=bounds_row["商品编码"])
        ws_bounds.cell(row=row_idx, column=2, value=bounds_row["商品名称"])
        value_col_idx = 3
        for ym_key in display_forecast_keys:
            for metric_key, _ in bounds_subheaders:
                raw_value = bounds_row.get(f"{ym_key}__{metric_key}", 0.0)
                cell = ws_bounds.cell(
                    row=row_idx,
                    column=value_col_idx,
                    value=raw_value if not (is_mom_row or is_yoy_row) else ("" if raw_value is None else raw_value),
                )
                cell.number_format = '0"%"' if (is_mom_row or is_yoy_row) else "0.0"
                value_col_idx += 1

        for col_idx in range(1, ws_bounds.max_column + 1):
            cell = ws_bounds.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if col_idx == 2 else "center", vertical="center")
            if is_total_row:
                cell.font = bold_font
                cell.fill = total_row_fill
            elif is_mom_row or is_yoy_row:
                cell.fill = grey_fill

    ws_bounds.column_dimensions["A"].width = 18
    ws_bounds.column_dimensions["B"].width = 42
    for col_idx in range(3, ws_bounds.max_column + 1):
        ws_bounds.column_dimensions[get_column_letter(col_idx)].width = 14

    ws = wb[DEMAND_SHEET_NAME]
    demand_header_to_col = {header: idx + 1 for idx, header in enumerate(final_cols)}
    demand_col_letter = {header: get_column_letter(col_idx) for header, col_idx in demand_header_to_col.items()}

    demand_row_map = {}
    group_ranges = []
    current_group_start = None
    for excel_row, raw_row in enumerate(all_results, start=2):
        row_code = raw_row.get("商品编码")
        if row_code not in {"合计", "环比", "同比"}:
            demand_row_map[row_code] = excel_row
            if current_group_start is None:
                current_group_start = excel_row
        elif row_code == "合计":
            group_ranges.append(
                {
                    "start": current_group_start,
                    "end": excel_row - 1,
                    "total": excel_row,
                    "mom": excel_row + 1,
                    "yoy": excel_row + 2,
                }
            )
            current_group_start = None

    const_turnover_ref = "'常量'!$B$4"
    forecast_col_letters = [demand_col_letter[key] for key in display_forecast_keys]
    reference_col_letters = [
        demand_col_letter[current_month_mean_col],
        demand_col_letter[current_month_window_col],
        demand_col_letter[current_month_ly_col],
    ]
    forecast_sum_col_letter = demand_col_letter["预测合计"]

    for code, excel_row in demand_row_map.items():
        raw_row = raw_demand_row(excel_row)
        ws_x_demand.write_formula(
            f"{forecast_sum_col_letter}{excel_row}",
            f"=SUM({forecast_col_letters[0]}{excel_row}:{forecast_col_letters[-1]}{excel_row})",
            None,
            raw_row.get("预测合计", 0),
        )
        for idx, forecast_key in enumerate(forecast_mom_keys, start=2):
            prev_col = demand_col_letter[display_forecast_keys[idx - 2]]
            current_col = demand_col_letter[forecast_key]
            mom_col = demand_col_letter[f"{forecast_key}环比(%)"]
            ws_x_demand.write_formula(
                f"{mom_col}{excel_row}",
                f'=IFERROR(({current_col}{excel_row}/{prev_col}{excel_row}-1)*100,"")',
                None,
                raw_row.get(f"{forecast_key}环比(%)", ""),
            )

    for group_range in group_ranges:
        start_row = group_range["start"]
        end_row = group_range["end"]
        total_row = group_range["total"]
        mom_row = group_range["mom"]
        yoy_row = group_range["yoy"]

        total_raw_row = raw_demand_row(total_row)
        mom_raw_row = raw_demand_row(mom_row)
        yoy_raw_row = raw_demand_row(yoy_row)

        for header in display_forecast_keys + [current_month_mean_col, current_month_window_col, current_month_ly_col]:
            col_letter = demand_col_letter[header]
            ws_x_demand.write_formula(
                f"{col_letter}{total_row}",
                f"=SUM({col_letter}{start_row}:{col_letter}{end_row})",
                None,
                total_raw_row.get(header, 0),
            )

        ws_x_demand.write_formula(
            f"{forecast_sum_col_letter}{total_row}",
            f"=SUM({forecast_col_letters[0]}{total_row}:{forecast_col_letters[-1]}{total_row})",
            None,
            total_raw_row.get("预测合计", 0),
        )

        for idx, forecast_key in enumerate(display_forecast_keys):
            col_letter = demand_col_letter[forecast_key]
            if idx == 0:
                prev_actual_col = demand_col_letter[current_year_actual_month_cols[-1]] if current_year_actual_month_cols else None
                if prev_actual_col:
                    ws_x_demand.write_formula(
                        f"{col_letter}{mom_row}",
                        f'=IFERROR(({col_letter}{total_row}/{prev_actual_col}{total_row}-1)*100,"")',
                        None,
                        mom_raw_row.get(forecast_key, ""),
                    )
                else:
                    ws[f"{col_letter}{mom_row}"] = ""
            else:
                prev_col = demand_col_letter[display_forecast_keys[idx - 1]]
                ws_x_demand.write_formula(
                    f"{col_letter}{mom_row}",
                    f'=IFERROR(({col_letter}{total_row}/{prev_col}{total_row}-1)*100,"")',
                    None,
                    mom_raw_row.get(forecast_key, ""),
                )

            yoy_actual_col = demand_col_letter.get(month_key(int(forecast_key.split("-")[0]) - 1, int(forecast_key.split("-")[1])))
            if yoy_actual_col:
                ws_x_demand.write_formula(
                    f"{col_letter}{yoy_row}",
                    f'=IFERROR(({col_letter}{total_row}/{yoy_actual_col}{total_row}-1)*100,"")',
                    None,
                    yoy_raw_row.get(forecast_key, ""),
                )

    pct_headers = set(f"{key}环比(%)" for key in forecast_mom_keys)

    ws.freeze_panes = ws["A2"]
    ws.row_dimensions[1].height = 42
    for ym_key, header_label in forecast_header_labels.items():
        ws[f"{demand_col_letter[ym_key]}1"] = header_label
    for cell in ws[1]:
        header = str(cell.value)
        if "参考" in header:
            cell.fill = yellow_fill
        elif "环比" in header:
            cell.fill = grey_fill
        elif any(token in header for token in ["实销合计", "最近12个月", "商品分类"]):
            cell.fill = blue_fill
        else:
            cell.fill = green_fill
        cell.font = note_font if "手动调整后会自动更新关联值" in header else header_font
        cell.alignment = header_align
        cell.border = border

    for row_index in range(2, ws.max_row + 1):
        raw_row = all_results[row_index - 2]
        first_cell_value = ws.cell(row=row_index, column=1).value
        is_total_row = first_cell_value == "合计"
        is_mom_row = first_cell_value == "环比"
        is_yoy_row = first_cell_value == "同比"
        max_12m = raw_row.get("最近12个月最高(不含本月)", 0)

        ws.row_dimensions[row_index].height = 22
        for col_index in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_index, column=col_index)
            col_name = final_cols[col_index - 1]
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if col_index == 2 else "center", vertical="center")

            if is_total_row:
                cell.font = bold_font
                cell.fill = total_row_fill
            elif is_mom_row or is_yoy_row:
                cell.fill = grey_fill

            raw_value = raw_row.get(col_name)
            if col_name in display_forecast_keys and not (is_mom_row or is_yoy_row):
                if raw_value and raw_value > max_12m:
                    cell.font = Font(color="FF0000", bold=is_total_row)

            is_pct_value = (col_name in pct_cols_in_output) or ((is_mom_row or is_yoy_row) and (col_name in month_cols_in_output))
            if col_name in pct_headers or ((is_mom_row or is_yoy_row) and (col_name in display_forecast_keys)):
                cell.number_format = '0"%"'
            if is_pct_value and raw_value is not None and not pd.isna(raw_value):
                if raw_value > 100 or raw_value < -100:
                    cell.font = Font(color="FF0000", bold=is_total_row)

    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min((max_len + 6) * 1.5, 55)

    ws_plan = wb["生产计划"]
    plan_header_to_col = {header: idx + 1 for idx, header in enumerate(plan_cols)}
    plan_col_letter = {header: get_column_letter(col_idx) for header, col_idx in plan_header_to_col.items()}
    ws_plan.freeze_panes = ws_plan["A2"]
    ws_plan.row_dimensions[1].height = 30
    plan_yellow_headers = {current_remaining_demand_col, *future_demand_cols, *future_daily_cols}
    plan_blue_headers = {current_gap_col, "库存可周转天数", "库存+在途可周转天数", "库存+在途是否满足需求", "库存计划量"}
    if cumulative_gap_col:
        plan_blue_headers.add(cumulative_gap_col)
    for cell in ws_plan[1]:
        header = str(cell.value)
        if header in plan_yellow_headers:
            cell.fill = yellow_fill
        elif header in plan_blue_headers:
            cell.fill = blue_fill
        else:
            cell.fill = green_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    current_daily_expr_by_row = {}
    last_labels = [f"{int(ym_key.split('-')[1])}月底以后" for ym_key in display_forecast_keys]
    period_days = [remaining_days_in_base_month] + [get_days_in_month(*map(int, ym_key.split("-"))) for ym_key in display_forecast_keys[1:]]

    for plan_row_idx, plan_row in enumerate(plan_results, start=2):
        code = plan_row["商品编码"]
        demand_row = demand_row_map[code]
        current_forecast_formula_expr = f"'{DEMAND_SHEET_NAME}'!${demand_col_letter[base_month_key]}${demand_row}"
        current_daily_expr = f"({current_forecast_formula_expr})/{days_in_base_month}"
        current_daily_expr_by_row[plan_row_idx] = current_daily_expr

        ws_x_plan.write_formula(
            f"{plan_col_letter[current_remaining_demand_col]}{plan_row_idx}",
            f"=ROUNDUP(({current_forecast_formula_expr})/{days_in_base_month}*{remaining_days_in_base_month},0)",
            None,
            plan_row.get(current_remaining_demand_col, 0),
        )

        for offset, demand_col in enumerate(future_demand_cols, start=2):
            forecast_key = display_forecast_keys[offset - 1]
            ws_x_plan.write_formula(
                f"{plan_col_letter[demand_col]}{plan_row_idx}",
                f"=ROUNDUP('{DEMAND_SHEET_NAME}'!${demand_col_letter[forecast_key]}${demand_row},0)",
                None,
                plan_row.get(demand_col, 0),
            )

        for offset, daily_col in enumerate(future_daily_cols, start=2):
            forecast_key = display_forecast_keys[offset - 1]
            month_days = get_days_in_month(*map(int, forecast_key.split("-")))
            ws_x_plan.write_formula(
                f"{plan_col_letter[daily_col]}{plan_row_idx}",
                f"='{DEMAND_SHEET_NAME}'!${demand_col_letter[forecast_key]}${demand_row}/{month_days}",
                None,
                plan_row.get(daily_col, 0),
            )

        inventory_total_expr = f"{plan_col_letter['当前在途数量']}{plan_row_idx}+{plan_col_letter['当前库存数量']}{plan_row_idx}"
        ws_x_plan.write_formula(
            f"{plan_col_letter[current_gap_col]}{plan_row_idx}",
            f"=MAX(0,{plan_col_letter[current_remaining_demand_col]}{plan_row_idx}-({inventory_total_expr}))",
            None,
            plan_row.get(current_gap_col, 0),
        )
        if cumulative_gap_col:
            ws_x_plan.write_formula(
                f"{plan_col_letter[cumulative_gap_col]}{plan_row_idx}",
                f'=MAX(0,{plan_col_letter[current_remaining_demand_col]}{plan_row_idx}+{plan_col_letter[future_demand_cols[0]]}{plan_row_idx}-({inventory_total_expr}))',
                None,
                plan_row.get(cumulative_gap_col, 0),
            )

        future_demand_refs = [f"{plan_col_letter[col]}{plan_row_idx}" for col in future_demand_cols]
        future_daily_refs = [f"{plan_col_letter[col]}{plan_row_idx}" for col in future_daily_cols]
        inventory_ref = f"{plan_col_letter['当前库存数量']}{plan_row_idx}"
        inventory_plus_transit_ref = f"({inventory_total_expr})"
        ws_x_plan.write_formula(
            f"{plan_col_letter['库存可周转天数']}{plan_row_idx}",
            build_turnover_formula(
                inventory_ref,
                f"{plan_col_letter[current_remaining_demand_col]}{plan_row_idx}",
                current_daily_expr,
                future_demand_refs,
                future_daily_refs,
                period_days,
                last_labels,
            ),
            None,
            plan_row.get("库存可周转天数", 0),
        )
        ws_x_plan.write_formula(
            f"{plan_col_letter['库存+在途可周转天数']}{plan_row_idx}",
            build_turnover_formula(
                inventory_plus_transit_ref,
                f"{plan_col_letter[current_remaining_demand_col]}{plan_row_idx}",
                current_daily_expr,
                future_demand_refs,
                future_daily_refs,
                period_days,
                last_labels,
            ),
            None,
            plan_row.get("库存+在途可周转天数", 0),
        )
        ws_x_plan.write_formula(
            f"{plan_col_letter['库存+在途是否满足需求']}{plan_row_idx}",
            f'=IF(OR(ISTEXT({plan_col_letter["库存+在途可周转天数"]}{plan_row_idx}),{plan_col_letter["库存+在途可周转天数"]}{plan_row_idx}>={const_turnover_ref}),"是","否")',
            None,
            plan_row.get("库存+在途是否满足需求", ""),
        )
        ws_x_plan.write_formula(
            f"{plan_col_letter['库存计划量']}{plan_row_idx}",
            build_inventory_plan_formula(
                inventory_plus_transit_ref,
                current_daily_expr,
                future_daily_refs,
                period_days,
            ),
            None,
            plan_row.get("库存计划量", 0),
        )

    for row_index in range(2, ws_plan.max_row + 1):
        ws_plan.row_dimensions[row_index].height = 22
        for col_index in range(1, ws_plan.max_column + 1):
            cell = ws_plan.cell(row=row_index, column=col_index)
            header = plan_cols[col_index - 1]
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if header in {"商品名称", "供应商"} else "center", vertical="center")
            if header == "库存+在途是否满足需求" and cell.value == "否":
                cell.font = Font(color="FF0000", bold=True)
            if header in {current_gap_col, cumulative_gap_col, "库存计划量"} and isinstance(cell.value, (int, float)) and cell.value > 0:
                cell.fill = red_fill

    for col in ws_plan.columns:
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col)
        ws_plan.column_dimensions[get_column_letter(col[0].column)].width = min((max_len + 6) * 1.5, 40)

    ws_constants = wb["常量"]
    ws_constants.freeze_panes = ws_constants["A2"]
    ws_constants.row_dimensions[1].height = 26
    for cell in ws_constants[1]:
        cell.fill = blue_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    for row_index in range(2, ws_constants.max_row + 1):
        note_text = str(ws_constants.cell(row=row_index, column=3).value or "")
        wrapped_line_count = max(1, (len(note_text) // 32) + 1)
        ws_constants.row_dimensions[row_index].height = max(22, min(110, 18 * wrapped_line_count))
        for col_index in range(1, ws_constants.max_column + 1):
            cell = ws_constants.cell(row=row_index, column=col_index)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if col_index == 1:
                cell.font = bold_font
                cell.fill = yellow_fill
            elif row_index in {4, 5}:
                cell.fill = green_fill
                if col_index == 3:
                    cell.font = note_font

    ws_constants.column_dimensions["A"].width = 32
    ws_constants.column_dimensions["B"].width = 28
    ws_constants.column_dimensions["C"].width = 72

    ws_formula = wb["公式"]
    ws_formula.freeze_panes = ws_formula["A2"]
    ws_formula.row_dimensions[1].height = 26
    for cell in ws_formula[1]:
        cell.fill = blue_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    for row_index in range(2, ws_formula.max_row + 1):
        ws_formula.row_dimensions[row_index].height = 42
        for col_index in range(1, ws_formula.max_column + 1):
            cell = ws_formula.cell(row=row_index, column=col_index)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if col_index == 1:
                cell.font = bold_font
                cell.fill = yellow_fill

    ws_formula.column_dimensions["A"].width = 28
    ws_formula.column_dimensions["B"].width = 120

    wb.save(file_path)

    # Inject cached formula values for Apple Numbers compatibility
    if demand_caches or plan_caches:
        sheet_formula_caches = {1: demand_caches, 2: plan_caches}
        inject_cached_values_for_numbers(file_path, sheet_formula_caches)

    print(f"✅ 导出成功: {file_path}")
    return file_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=f"导出 {DEMAND_SHEET_NAME} 与生产计划工作簿")
    parser.add_argument(
        "--sku-scope-file",
        help="SKU 范围文件路径；文件中需包含商品编码列或兼容列名，仅导出该范围内的 SKU",
    )
    parser.add_argument(
        "--months",
        type=int,
        default=DEFAULT_FORECAST_MONTHS,
        help=f"AI 预测导出月数，默认 {DEFAULT_FORECAST_MONTHS}，最多不超过 jiuyan_forecasts.json 中的 months 数量",
    )
    parser.add_argument(
        "--forecast-json",
        default=FORECAST_JSON_PATH,
        help=f"预测结果 JSON 路径，默认 {FORECAST_JSON_PATH}",
    )
    args = parser.parse_args()
    export_forecast(
        sku_scope_file=args.sku_scope_file,
        forecast_json_path=args.forecast_json,
        forecast_months=args.months,
    )
